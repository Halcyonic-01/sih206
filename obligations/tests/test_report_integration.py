"""P5 acceptance through the default K1 routes, with labeled synthetic model rows."""
from datetime import date, datetime, time, timedelta
from io import BytesIO, StringIO
import json
from pathlib import Path
from time import perf_counter
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.core.management import call_command
from django.db import OperationalError
from django.test import Client, override_settings
from django.utils import timezone
from openpyxl import load_workbook

from access.models import AccessLog
from obligations import reports
from obligations.catalogue import draft_payload
from obligations.models import CsvImport, DepositObligation, MasterAudit, RuleSet
from portal.models import AuditEvent, Document, Institution, Instrument
from portal.tests.common import PortalTestCase
from verification.models import VerificationEvidence


@override_settings(ROOT_URLCONF='config.urls')
class ReportIntegrationTests(PortalTestCase):
    def setUp(self):
        self.client.force_login(self.officer)
        self.when = date(2026, 10, 25)

    def item(self, **values):
        base = dict(institution=self.institution_a, created_by=self.a,
            reference='SYNTH-' + str(Instrument.objects.count()), bank='Synthetic bank',
            amount_paise=10000, issue_date=date(2026, 1, 1), maturity_date=date(2026, 11, 1),
            status='ACCEPTED')
        return Instrument.objects.create(**(base | values))

    def rows(self, kind, actor=None):
        return reports.report_rows(actor or self.officer, kind, as_of=self.when)['rows']

    def evidence(self, item, method='BANK_CONFIRMATION', result='CONFIRMED', **values):
        return VerificationEvidence.objects.create(instrument=item, actor=self.officer,
            instrument_version=item.version, method=method, result=result,
            source='SYNTHETIC TEST EVIDENCE; no bank was contacted', **values)

    def post(self, path, data):
        return self.client.post('/api/' + path, json.dumps(data), content_type='application/json')

    def preview(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        content = (Path(__file__).resolve().parents[2] / 'evaluation/s4/valid.csv').read_bytes()
        result = self.client.post('/api/imports/institutions/preview',
            {'file': SimpleUploadedFile('SIMULATED.csv', content, content_type='text/csv')})
        self.assertEqual(result.status_code, 200, result.content)
        self.assertEqual(result.json()['simulation_label'], 'SIMULATED')
        return result.json()['preview_id']

    def test_five_report_memberships_match_labeled_snapshot_and_scope(self):
        pending = self.item(status='SUBMITTED')
        confirmed = self.item(verification_level='CONFIRMED')
        self.evidence(confirmed)
        simulated = self.item(verification_level='REVIEWED')
        self.evidence(simulated, method='SIMULATED_STATEMENT', result='SIMULATED_MATCH')
        released = self.item(released_at=timezone.now())
        old = self.item(superseded_at=timezone.now())
        foreign = self.item(institution=self.institution_b, created_by=self.b)
        draft = self.item(status='DRAFT')
        pending_ids = {r['instrument_id'] for r in self.rows('pending_verifications')}
        self.assertEqual(pending_ids, {pending.pk, simulated.pk, foreign.pk})
        self.assertEqual({r['instrument_id'] for r in self.rows('upcoming_maturities')},
            {confirmed.pk, simulated.pk, foreign.pk})
        self.assertEqual({r['instrument_id'] for r in self.rows('released_register')}, {released.pk})
        self.assertNotIn(draft.pk, {r['instrument_id'] for r in self.rows('missing_documents')})
        self.assertNotIn(old.pk, pending_ids)
        self.assertEqual(self.rows('pending_verifications', self.b)[0]['instrument_id'], foreign.pk)
        self.assertEqual(self.rows('released_register', self.b), [])
        self.assertEqual(self.rows('intervention'), [])

    def test_current_confirmation_requires_actual_method_and_version(self):
        stale = self.item(version=2, verification_level='CONFIRMED')
        VerificationEvidence.objects.create(instrument=stale, actor=self.officer, instrument_version=1,
            method='BANK_CONFIRMATION', result='CONFIRMED', source='SYNTHETIC stale')
        unproven = self.item(verification_level='CONFIRMED')
        self.assertEqual({r['instrument_id'] for r in self.rows('pending_verifications')}, {stale.pk, unproven.pk})

    def test_document_types_follow_each_deposit_mode(self):
        rows = [self.item(mode=mode) for mode in ['FDR', 'RTGS', 'RPGF']]
        kinds = ['FDR', 'PAYMENT_PROOF', 'RPGF_PROOF']
        for item, kind in zip(rows, kinds):
            self.assertIn(item.pk, {r['instrument_id'] for r in self.rows('missing_documents')})
            Document.objects.create(instrument=item, uploaded_by=self.a, kind=kind, file='synthetic-not-read.pdf',
                original_name='SYNTHETIC.pdf', size_bytes=1, sha256='0'*64, page_count=1)
        self.assertEqual(self.rows('missing_documents'), [])

    def test_sla_exact_boundaries_and_actual_config(self):
        submitted = self.item(status='SUBMITTED')
        changed = self.item(status='CHANGES_REQUESTED')
        for item, action, day in [(submitted, 'SUBMITTED', date(2026, 10, 22)),
                                  (changed, 'CHANGES_REQUESTED', date(2026, 10, 10))]:
            event = AuditEvent.objects.create(instrument=item, actor=self.a, action=action)
            at = timezone.make_aware(datetime.combine(day, time(12)))
            AuditEvent.objects.filter(pk=event.pk).update(created_at=at)
        RuleSet.objects.create(code='escalation-matrix', version=1, effective_from=date(2026, 1, 1),
            source_document='Synthetic', source_clause='Test threshold',
            payload={'sla': {'submission_working_days': 1, 'correction_days': [15, 30]}})
        # Fri is exactly 1 working day after Thursday: not breached. Sunday still 1.
        result = self.rows('intervention')
        self.assertEqual([r['instrument_id'] for r in result], [changed.pk])
        self.assertIn('Corrections ageing', result[0]['reason'])
        next_day = reports.report_rows(self.officer, 'intervention', as_of=date(2026, 10, 26))
        self.assertEqual({r['instrument_id'] for r in next_day['rows']}, {submitted.pk, changed.pk})

    def test_intervention_has_real_discrepancy_and_unlinked_obligation_shortfall(self):
        item = self.item(verification_level='DISCREPANCY', maturity_date=date(2026, 10, 24))
        rule = RuleSet.objects.create(code='synth-report', version=1, effective_from=date(2026, 1, 1),
            source_document='Synthetic', source_clause='Draft', payload=draft_payload())
        obligation = DepositObligation.objects.create(institution=self.institution_a, required_paise=10000,
            term_start=date(2026, 1, 1), term_end=date(2036, 1, 1), rule_version=rule)
        rows = self.rows('intervention')
        self.assertEqual(len(rows), 2)
        self.assertIn('discrepancy', rows[0]['reason'])
        self.assertIn('Renewal overdue', rows[0]['reason'])
        self.assertEqual(rows[1]['reference'], f'Obligation #{obligation.pk}')
        self.assertIsNone(rows[1]['instrument_id'])
        self.assertEqual(self.rows('intervention', self.b), [])

    def test_reports_log_access_fail_closed_and_bound_exports(self):
        self.item()
        result = self.client.get('/api/reports/pending_verifications')
        self.assertEqual(result.status_code, 200)
        self.assertEqual(AccessLog.objects.filter(action='REPORT', object_type='s4_report').count(), 1)
        with patch('obligations.api.report_views.log_access', side_effect=OperationalError('private logging failure')):
            result = self.client.get('/api/reports/pending_verifications?download=xlsx')
        self.assertEqual(result.status_code, 503)
        self.assertNotIn(b'private logging failure', result.content)
        with patch('obligations.report_sources.MAX_RECORDS', 0):
            self.assertEqual(self.client.get('/api/reports/pending_verifications').status_code, 400)

    def test_report_invalid_auth_and_registered_p6(self):
        self.assertEqual(Client().get('/api/reports/pending_verifications').status_code, 401)
        for path in ['reports/unknown', 'reports/intervention?institution_id=2', 'reports/intervention?download=csv']:
            self.assertEqual(self.client.get('/api/' + path).status_code, 400)
        self.assertEqual(self.client.get('/api/releases').status_code, 200)
        group = Group.objects.create(name='Clearance')
        user = get_user_model().objects.create_user(username='SYNTH-CLEARANCE')
        user.groups.add(group)
        self.client.force_login(user)
        self.assertEqual(self.client.get('/api/reports/missing_documents').status_code, 403)

    def test_csv_no_writes_until_commit_replay_and_foreign_actor(self):
        before = Institution.objects.count()
        pk = self.preview()
        self.assertEqual(Institution.objects.count(), before)
        officer = get_user_model().objects.create_user(username='SYNTH-OTHER-OFFICER')
        officer.groups.add(Group.objects.get(name='Officer'))
        self.client.force_login(officer)
        payload = {'version': 1, 'reason': 'SIMULATED complete import'}
        self.assertEqual(self.post(f'imports/institutions/{pk}/commit', payload).status_code, 404)
        self.client.force_login(self.officer)
        path = f'imports/institutions/{pk}/commit'
        self.assertEqual(self.post(path, {'version': True, 'reason': 'Invalid version'}).status_code, 400)
        result = self.post(path, payload)
        self.assertEqual(result.status_code, 200, result.content)
        self.assertEqual(Institution.objects.count(), before + 3)
        self.assertEqual(self.post(path, payload).status_code, 409)
        self.assertEqual(CsvImport.objects.get(pk=pk).source, b'')

    def test_csv_hash_expiry_and_rollback_preserve_all_rows(self):
        pk = self.preview()
        CsvImport.objects.filter(pk=pk).update(source_sha256='0'*64)
        self.assertEqual(self.post(f'imports/institutions/{pk}/commit', {'version':1,'reason':'Test'}).status_code,409)
        pk = self.preview()
        CsvImport.objects.filter(pk=pk).update(expires_at=timezone.now()-timedelta(seconds=1))
        self.assertEqual(self.post(f'imports/institutions/{pk}/commit', {'version':1,'reason':'Test'}).status_code,409)
        pk = self.preview()
        before = Institution.objects.count()
        from obligations.imports import create_institution
        calls = 0
        def failure(*args, **kwargs):
            nonlocal calls
            calls += 1
            if calls == 2: raise OperationalError('synthetic DB failure')
            return create_institution(*args, **kwargs)
        with patch('obligations.imports.create_institution', side_effect=failure):
            self.assertEqual(self.post(f'imports/institutions/{pk}/commit', {'version':1,'reason':'Test'}).status_code,503)
        self.assertEqual(Institution.objects.count(), before)
        self.assertEqual(MasterAudit.objects.count(), 0)
        self.assertIsNone(CsvImport.objects.get(pk=pk).committed_at)

    def test_import_invalid_file_role_and_csrf(self):
        self.assertEqual(self.client.post('/api/imports/institutions/preview', {}).status_code,400)
        self.client.force_login(self.a)
        self.assertEqual(self.client.post('/api/imports/institutions/preview', {}).status_code,403)
        client = Client(enforce_csrf_checks=True)
        client.force_login(self.officer)
        self.assertEqual(client.post('/api/imports/institutions/preview', {}).status_code,403)

    def test_xlsx_is_exact_formula_safe_and_measured_at_200_rows(self):
        amount = 9007199254740993
        for n in range(200): self.item(reference='=SYNTH-' + str(n), amount_paise=amount if n%2==0 else amount-2)
        start = perf_counter()
        result = reports.report_rows(self.officer, 'pending_verifications', as_of=self.when)
        raw = reports.xlsx_bytes(result)
        elapsed = perf_counter()-start
        sheet = load_workbook(BytesIO(raw)).active
        header = next(row[0].row for row in sheet if row[0].value == 'instrument_id')
        self.assertEqual(sheet.max_row-header, 200)
        self.assertEqual(sheet.cell(header+1,3).data_type, 's')
        self.assertEqual(sheet.cell(header+1,4).value, str(amount))
        self.assertEqual(sheet.cell(header+2,4).value, str(amount-2))
        self.assertEqual(sheet.auto_filter.ref, f'A{header}:H{sheet.max_row}')
        print(f'P5 bounded export: 200 rows, {len(raw)} bytes, {elapsed:.3f} seconds (synthetic local test).')
        response = self.client.get('/api/reports/pending_verifications', {'page_size': 10, 'page': 2})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['count'], 200)
        self.assertEqual(len(response.json()['results']), 10)
        self.assertEqual(response.json()['rows'], response.json()['results'])

    def test_seed_is_discovered_idempotent_and_preserves_existing_models(self):
        from django.core.management import get_commands
        self.assertIn('seed_review_obligations',get_commands())
        call_command('seed_review_obligations',stdout=StringIO())
        before = (Institution.objects.count(),DepositObligation.objects.count(),MasterAudit.objects.count())
        call_command('seed_review_obligations',stdout=StringIO())
        self.assertEqual(before,(Institution.objects.count(),DepositObligation.objects.count(),MasterAudit.objects.count()))
        self.assertEqual(Instrument.objects.count(),0)

    def test_kpi_cohort_real_evidence_not_simulated_and_incomplete_history_unknown(self):
        now = timezone.now()
        base = now-timedelta(days=3)
        known = self.item(verification_level='CONFIRMED')
        simulated = self.item(verification_level='REVIEWED')
        incomplete = self.item()
        for item in [known,simulated]:
            for n,action in enumerate(['DRAFT_CREATED','SUBMITTED','CHANGES_REQUESTED','SUBMITTED','ACCEPTED']):
                event=AuditEvent.objects.create(instrument=item,actor=self.a,action=action,
                    after={'verification_level':'DECLARED'} if n==0 else {})
                AuditEvent.objects.filter(pk=event.pk).update(created_at=base+timedelta(hours=n))
        self.evidence(known)
        self.evidence(simulated,method='SIMULATED_BANK',result='SIMULATED_MATCH')
        result=reports.kpi_report(self.officer,window_start=base,window_end=now,as_of=timezone.now())
        self.assertEqual(result['submission_to_acceptance']['sample_count'],2)
        self.assertEqual(result['submission_to_acceptance']['median_seconds'],'10800.000000')
        self.assertEqual(result['correction_cycles']['total'],2)
        # The simulated bank-link response does not confirm, but that record's history is complete, so it is
        # known to be unconfirmed rather than unknown: 1 of 2 confirmed.
        self.assertEqual(result['confirmed_share']['percentage_bps'],5000)
        self.assertEqual(result['confirmed_share']['numerator'],1)
        self.assertEqual(result['confirmed_share']['unknown_count'],0)
        self.assertEqual(result['data_quality']['incomplete_process_histories'],1)
        self.assertEqual(reports.kpi_report(self.b,window_start=base,window_end=now,as_of=timezone.now())['data_quality']['process_snapshots'],0)

    def test_kpi_api_invalid_windows_logging_and_database_failure(self):
        end=timezone.now()
        query={'window_start':(end-timedelta(days=1)).isoformat(),'window_end':end.isoformat()}
        result=self.client.get('/api/kpis',query)
        self.assertEqual(result.status_code,200,result.content)
        self.assertTrue(AccessLog.objects.filter(object_type='s4_kpis').exists())
        for values in [{},{'window_start':'bad','window_end':'bad'},
            query|{'window_start':'2026-01-01'},query|{'window_start':(end+timedelta(days=1)).isoformat()}]:
            self.assertEqual(self.client.get('/api/kpis',values).status_code,400)
        with patch('obligations.report_sources.records_for',side_effect=OperationalError('private SQL')):
            result=self.client.get('/api/kpis',query)
        self.assertEqual(result.status_code,503)
        self.assertNotIn(b'private SQL',result.content)
