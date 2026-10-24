"""S4 scoped read adapters. No fake verification or notification success."""
from datetime import timedelta
from io import BytesIO

from django.core.exceptions import ValidationError
from django.utils import timezone

from portal.models import AuditEvent
from portal.permissions import scoped_instruments
from .kpis import AcceptedSuccessor, KpiEvent, ProcessSnapshot, RenewalSnapshot, calculate_kpis
from .services import wire_money
from . import report_sources as sources

REPORTS = {'pending_verifications', 'upcoming_maturities', 'missing_documents', 'intervention', 'released_register'}


def report_rows(actor, kind, *, as_of=None):
    if kind not in REPORTS:
        raise ValidationError({'kind': 'Unsupported report.'})
    when = as_of or timezone.localdate()
    records = sources.records_for(actor)
    limitations = ['Stored typed documents are not reviewed receipt dates or proof of authenticity.',
        'Release-path checklists await P6/P7; this report checks deposit-mode supporting documents only.',
        'Current persisted state, not historical reconstruction. Draft policy assessments are advisory.']
    sla = sources.sla_context(when) if kind == 'intervention' else None
    rows = []
    for h in records:
        documents = list(h.documents.all())
        reasons = []
        if kind == 'pending_verifications':
            if h.status == 'DRAFT' or h.released_at or h.superseded_at or sources.verified_current(h):
                continue
            reasons = ['Independent confirmation pending; level=' + h.verification_level]
        elif kind == 'upcoming_maturities':
            if not h.is_current or not when <= h.maturity_date <= when + timedelta(days=90):
                continue
            reasons = ['Current holding matures within 90 days']
        elif kind == 'missing_documents':
            required = {'FDR': 'FDR', 'RTGS': 'PAYMENT_PROOF', 'RPGF': 'RPGF_PROOF'}[h.mode]
            if any(d.kind == required for d in documents):
                continue
            reasons = [required + ' supporting document missing']
        elif kind == 'intervention':
            if h.is_current and h.maturity_date < when:
                reasons.append('Renewal overdue')
            if any(o.severity == 'BLOCKING' and o.resolved_at is None for o in h.observations.all()):
                reasons.append('Unresolved blocking observation')
            if h.verification_level == 'DISCREPANCY':
                reasons.append('Verification discrepancy requiring human review')
            from notifications.rules import working_days_after
            if h.status == 'SUBMITTED':
                since = sources.latest_action(h, 'SUBMITTED', when)
                if since and working_days_after(since, when, sla[2]) > sla[0]:
                    reasons.append('Submission SLA exceeded; ' + sla[3])
                elif since is None:
                    limitations.append('Submission timestamp unavailable for instrument ' + str(h.pk))
            if h.status == 'CHANGES_REQUESTED':
                since = sources.latest_action(h, 'CHANGES_REQUESTED', when)
                if since and (when - since).days >= sla[1]:
                    reasons.append('Corrections ageing; ' + sla[3])
                elif since is None:
                    limitations.append('Correction timestamp unavailable for instrument ' + str(h.pk))
            if not reasons:
                continue
        elif h.released_at is None:
            continue
        rows.append({'instrument_id': h.pk, 'institution': h.institution.name, 'reference': h.reference,
            'amount_paise': h.amount_paise, 'maturity_date': h.maturity_date.isoformat(),
            'status': h.status, 'released_at': h.released_at.isoformat() if h.released_at else None,
            'reason': '; '.join(reasons)})
    if kind == 'intervention':
        from .services import scoped_obligations, obligation_detail
        obligations = list(scoped_obligations(actor).select_related('institution').order_by('pk')[:sources.MAX_RECORDS + 1])
        if len(obligations) > sources.MAX_RECORDS:
            raise ValidationError('At most 5000 visible obligations are supported. No partial export was produced.')
        for ob in obligations:
            assessment = obligation_detail(actor, ob.pk, as_of=when)['assessment']
            if assessment['status'] != 'AVAILABLE':
                limitations.append('Obligation assessment unavailable: ' + str(ob.pk))
                continue
            reasons = [f['explanation'] for f in assessment['findings'] if f['code'] in {'AMOUNT_SHORTFALL', 'TERM_GAP'}]
            if reasons:
                rows.append({'instrument_id': None, 'institution': ob.institution.name,
                    'reference': f'Obligation #{ob.pk}', 'amount_paise': ob.required_paise,
                    'maturity_date': ob.term_end.isoformat() if ob.term_end else None,
                    'status': ob.status, 'released_at': None, 'reason': '; '.join(reasons)})
    return wire_money({'kind': kind, 'as_of': when.isoformat(), 'simulation_label': 'SIMULATED',
        'limitations': list(dict.fromkeys(limitations)), 'rows': rows})


def xlsx_bytes(report):
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'SIMULATED report'
    fields = ['instrument_id', 'institution', 'reference', 'amount_paise', 'maturity_date', 'status', 'released_at', 'reason']
    sheet.append(['SIMULATED — not independently bank-verified'])
    sheet.append(['Report', report['kind'], 'As of', report['as_of']])
    for limitation in report['limitations']:
        sheet.append(['Limitation', limitation])
    sheet.append(fields)
    header = sheet.max_row
    for row in report['rows']:
        # Explicit string cell type prevents formula injection, including leading whitespace.
        # Excel numeric cells retain only 15 significant digits, less than JS's
        # safe-integer range. Store paise as decimal text even inside that range.
        sheet.append([str(row[key]) if key == 'amount_paise' and row[key] is not None else
                      ILLEGAL_CHARACTERS_RE.sub('\ufffd', row[key]) if isinstance(row[key], str)
                      else row[key] for key in fields])
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = 's'
    # All metadata is also untrusted text, including dynamically supplied limitations.
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                cell.value = ILLEGAL_CHARACTERS_RE.sub('\ufffd', cell.value)
                cell.data_type = 's'
    sheet.freeze_panes = f'A{header + 1}'
    sheet.auto_filter.ref = f'A{header}:H{max(header, sheet.max_row)}'
    for column in ['A', 'D', 'F']:
        sheet.column_dimensions[column].width = 22
    for column in ['B', 'C', 'E', 'G', 'H']:
        sheet.column_dimensions[column].width = 38
    target = BytesIO()
    workbook.save(target)
    return target.getvalue()


def renewal_required(holding):
    obligation = holding.obligation
    if obligation is None or obligation.term_end is None:
        return None
    return holding.maturity_date < obligation.term_end


def kpi_report(actor, *, window_start, window_end, as_of=None):
    cutoff = as_of or timezone.now()
    holdings = sources.records_for(actor)
    histories = {}
    for h in holdings:
        histories[h.pk] = sorted(h.audit_events.all(), key=lambda e: (e.created_at, e.pk))
    processes, renewals = [], []
    event_identity = 0
    for h in holdings:
        history = histories[h.pk]
        complete = bool(history and history[0].action in {'DRAFT_CREATED', 'RENEWAL_DRAFT_CREATED'})
        review = [e for e in history if e.action in {'SUBMITTED', 'CHANGES_REQUESTED', 'ACCEPTED'}]
        # Only actual non-simulated bank evidence supports confirmation metrics.
        # Missing/stale evidence is unknown; a simulated match is never confirmation.
        evidence = sorted((e for e in h.verification_evidence.all() if sources.independent(e)
            and e.instrument_version == h.version), key=lambda e: (e.created_at, e.pk))
        combined = [(e.created_at, e.pk, e.action) for e in review]
        if complete and evidence and history[0].after.get('verification_level') == 'DECLARED':
            combined.extend([(history[0].created_at, -1, 'DECLARED'),
                             (evidence[0].created_at, -2, 'INDEPENDENTLY_CONFIRMED')])
        combined.sort(key=lambda e: (e[0], 0 if e[2] == 'DECLARED' else
            2 if e[2] == 'INDEPENDENTLY_CONFIRMED' else 1, e[1]))
        # Local sequence identities are unique across process snapshots; source PKs
        # from two tables are not assumed to share a primary-key namespace.
        events = tuple(KpiEvent(event_identity + sequence, sequence, action, at)
            for sequence, (at, _, action) in enumerate(combined, 1))
        event_identity += len(events)
        processes.append(ProcessSnapshot(h.pk, h.institution_id, complete, events))
        # Acceptance and renewal eligibility are distinct facts.
        accepted = any(e.action == 'ACCEPTED' and e.created_at <= cutoff for e in history)
        if not accepted:
            continue
        successors = [s for s in holdings if s.predecessor_id == h.pk]
        successor = None
        if successors:
            candidate = successors[0]
            accepted_events = [e for e in histories[candidate.pk] if e.action == 'ACCEPTED']
            if accepted_events:
                successor = AcceptedSuccessor(candidate.pk, candidate.institution_id, accepted_events[0].created_at)
        renewals.append(RenewalSnapshot(h.pk, h.institution_id, h.maturity_date,
            # The obligation defines whether renewal is required, not acceptance alone: a holding that matures
            # before its obligation's term ends must be renewed (APH 10-year term). No linked term stays unknown.
            renewal_required(h), complete, successor))
    result = calculate_kpis(processes=tuple(processes), renewals=tuple(renewals),
        scope_institution_ids=tuple(sorted({h.institution_id for h in holdings})),
        window_start=window_start, window_end=window_end, as_of=cutoff, simulation_label='SIMULATED')
    result['limitations'] = ['Renewal is counted as required when a holding matures before its linked obligation term ends; '
        'holdings without a linked obligation term are excluded as unknown.',
        'Only current-version BANK_CONFIRMATION/CONFIRMED evidence counts; simulated matches never count.',
        'Missing creation/declaration history remains unknown. No before/after improvement study is implied.']
    return result
